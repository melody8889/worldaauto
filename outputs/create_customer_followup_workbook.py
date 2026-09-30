from pathlib import Path
import shutil
import zipfile

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation


DESKTOP = Path(r"C:\Users\Melody\Desktop")
SRC = DESKTOP / "汽车外贸客户跟进管理表_分客户专属sheet_豆包AI生成.xlsx"
OUT_XLSX = DESKTOP / "汽车外贸客户跟进管理表_分客户专属sheet_整理版.xlsx"
OUT_BAS = DESKTOP / "汽车外贸客户跟进管理表_VBA宏代码.bas"

SHEET_OVERVIEW = "客户总览"
SHEET_TEMPLATE = "跟进模板"

VBA_CODE = r'''Attribute VB_Name = "客户跟进自动化"
Option Explicit

Const SHEET_OVERVIEW As String = "客户总览"
Const SHEET_TEMPLATE As String = "跟进模板"

' 初始化固定的客户总览与跟进模板结构。
Public Sub 初始化客户跟进表()
    Dim wsO As Worksheet
    Dim wsT As Worksheet
    
    Set wsO = GetOrCreateSheet(SHEET_OVERVIEW)
    Set wsT = GetOrCreateSheet(SHEET_TEMPLATE)
    
    Application.ScreenUpdating = False
    
    wsO.Cells.Clear
    wsO.Range("A1:J1").Value = Array("序号", "客户公司", "联系人", "国家", "WhatsApp", "邮箱", "意向车型", "客户等级", "跳转跟进页", "最后跟进日期")
    
    With wsO.Range("A1:J1")
        .Font.Bold = True
        .Font.Color = vbWhite
        .Interior.Color = RGB(31, 78, 121)
        .HorizontalAlignment = xlCenter
        .VerticalAlignment = xlCenter
        .Borders.LineStyle = xlContinuous
    End With
    
    wsO.Rows(1).RowHeight = 28
    wsO.Columns("A").ColumnWidth = 8
    wsO.Columns("B:C").ColumnWidth = 24
    wsO.Columns("D:F").ColumnWidth = 18
    wsO.Columns("G:I").ColumnWidth = 24
    wsO.Columns("J").ColumnWidth = 18
    wsO.Range("A2:J1000").Borders.LineStyle = xlContinuous
    
    With wsO.Range("H2:H1000").Validation
        .Delete
        .Add Type:=xlValidateList, AlertStyle:=xlValidAlertStop, Operator:=xlBetween, Formula1:="A高意向,B意向,C潜在,D沉睡"
        .IgnoreBlank = True
        .InCellDropdown = True
    End With
    
    With wsO.Range("G2:G1000").Validation
        .Delete
        .Add Type:=xlValidateList, AlertStyle:=xlValidAlertStop, Operator:=xlBetween, Formula1:="Deepal S07,Deepal S05,Deepal SL03,ROX ADAMAS"
        .IgnoreBlank = True
        .InCellDropdown = True
    End With
    
    wsT.Cells.Clear
    wsT.Range("A1").Value = "客户公司："
    wsT.Range("A2").Value = "对接人："
    wsT.Range("A3").Value = "目标车型："
    wsT.Range("A4").Value = "国家："
    wsT.Range("A5").Value = "客户等级："
    wsT.Range("A7:E7").Value = Array("跟进日期", "跟进渠道", "跟进内容详情", "客户反馈", "下次跟进计划")
    
    With wsT.Range("A1:F5")
        .Font.Bold = True
        .Interior.Color = RGB(242, 242, 242)
        .Borders.LineStyle = xlContinuous
        .VerticalAlignment = xlCenter
    End With
    
    With wsT.Range("A7:E7")
        .Font.Bold = True
        .Font.Color = vbWhite
        .Interior.Color = RGB(166, 124, 82)
        .HorizontalAlignment = xlCenter
        .VerticalAlignment = xlCenter
        .Borders.LineStyle = xlContinuous
    End With
    
    wsT.Columns("A:B").ColumnWidth = 20
    wsT.Columns("C:E").ColumnWidth = 34
    wsT.Rows(7).RowHeight = 26
    wsT.Range("A8:E1000").Borders.LineStyle = xlContinuous
    
    With wsT.Range("B8:B1000").Validation
        .Delete
        .Add Type:=xlValidateList, AlertStyle:=xlValidAlertStop, Operator:=xlBetween, Formula1:="WhatsApp,邮件,电话,IG/Facebook社媒"
        .IgnoreBlank = True
        .InCellDropdown = True
    End With
    
    Application.ScreenUpdating = True
    MsgBox "客户跟进表初始化完成。", vbInformation
End Sub

' 批量读取客户总览，自动生成每个客户的专属跟进页并创建跳转链接。
Public Sub 一键生成客户跟进页()
    Dim wsO As Worksheet
    Dim wsT As Worksheet
    Dim wsNew As Worksheet
    Dim lastRow As Long
    Dim r As Long
    Dim companyName As String
    Dim contactName As String
    Dim sheetName As String
    Dim skipped As String
    
    Set wsO = Worksheets(SHEET_OVERVIEW)
    Set wsT = Worksheets(SHEET_TEMPLATE)
    
    lastRow = wsO.Cells(wsO.Rows.Count, "B").End(xlUp).Row
    Application.ScreenUpdating = False
    
    For r = 2 To lastRow
        companyName = Trim(CStr(wsO.Cells(r, "B").Value))
        contactName = Trim(CStr(wsO.Cells(r, "C").Value))
        
        If companyName <> "" Or contactName <> "" Then
            sheetName = companyName & "-" & contactName
            
            If Not IsValidSheetName(sheetName) Then
                skipped = skipped & vbCrLf & "第 " & r & " 行：" & sheetName
                GoTo NextRow
            End If
            
            If Not SheetExists(sheetName) Then
                wsT.Copy After:=Worksheets(Worksheets.Count)
                Set wsNew = ActiveSheet
                wsNew.Name = sheetName
            Else
                Set wsNew = Worksheets(sheetName)
            End If
            
            wsNew.Range("B1").Value = wsO.Cells(r, "B").Value
            wsNew.Range("B2").Value = wsO.Cells(r, "C").Value
            wsNew.Range("B3").Value = wsO.Cells(r, "G").Value
            wsNew.Range("B4").Value = wsO.Cells(r, "D").Value
            wsNew.Range("B5").Value = wsO.Cells(r, "H").Value
            
            wsO.Cells(r, "I").Hyperlinks.Delete
            wsO.Hyperlinks.Add Anchor:=wsO.Cells(r, "I"), Address:="", SubAddress:="'" & sheetName & "'!A1", TextToDisplay:="进入跟进页"
        End If
NextRow:
    Next r
    
    Application.ScreenUpdating = True
    
    If skipped <> "" Then
        MsgBox "以下客户名称含非法字符或长度超过31字符，已跳过：" & skipped, vbExclamation
    End If
    
    MsgBox "全部客户跟进页创建完成，链接已自动生成。", vbInformation
End Sub

' 删除所有客户专属sheet，只保留客户总览与跟进模板。
Public Sub 一键清空客户跟进页()
    Dim ws As Worksheet
    Dim i As Long
    
    Application.ScreenUpdating = False
    Application.DisplayAlerts = False
    
    For i = Worksheets.Count To 1 Step -1
        Set ws = Worksheets(i)
        If ws.Name <> SHEET_OVERVIEW And ws.Name <> SHEET_TEMPLATE Then
            ws.Delete
        End If
    Next i
    
    Application.DisplayAlerts = True
    Worksheets(SHEET_OVERVIEW).Range("I2:I1000").ClearContents
    Application.ScreenUpdating = True
    
    MsgBox "所有客户跟进页已清空，仅保留客户总览和跟进模板。", vbInformation
End Sub

Private Function GetOrCreateSheet(ByVal sheetName As String) As Worksheet
    If SheetExists(sheetName) Then
        Set GetOrCreateSheet = Worksheets(sheetName)
    Else
        Set GetOrCreateSheet = Worksheets.Add(After:=Worksheets(Worksheets.Count))
        GetOrCreateSheet.Name = sheetName
    End If
End Function

Private Function SheetExists(ByVal sheetName As String) As Boolean
    Dim ws As Worksheet
    SheetExists = False
    For Each ws In Worksheets
        If ws.Name = sheetName Then
            SheetExists = True
            Exit Function
        End If
    Next ws
End Function

Private Function IsValidSheetName(ByVal sheetName As String) As Boolean
    Dim illegalChars As Variant
    Dim ch As Variant
    
    illegalChars = Array("/", "\", "?", "*", "[", "]", ":")
    
    If Len(sheetName) = 0 Or Len(sheetName) > 31 Then
        IsValidSheetName = False
        Exit Function
    End If
    
    For Each ch In illegalChars
        If InStr(sheetName, ch) > 0 Then
            IsValidSheetName = False
            Exit Function
        End If
    Next ch
    
    IsValidSheetName = True
End Function
'''


def reset_workbook() -> Workbook:
    wb = Workbook()
    ws_o = wb.active
    ws_o.title = SHEET_OVERVIEW
    ws_t = wb.create_sheet(SHEET_TEMPLATE)
    return wb


def style_workbook(wb: Workbook) -> None:
    thin = Side(style="thin", color="D9DEE7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    blue = PatternFill("solid", fgColor="1F4E79")
    gold = PatternFill("solid", fgColor="A67C52")
    pale = PatternFill("solid", fgColor="F2F2F2")

    ws_o = wb[SHEET_OVERVIEW]
    headers = ["序号", "客户公司", "联系人", "国家", "WhatsApp", "邮箱", "意向车型", "客户等级", "跳转跟进页", "最后跟进日期"]
    ws_o.append(headers)
    for cell in ws_o[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = blue
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    ws_o.row_dimensions[1].height = 28
    widths = {"A": 8, "B": 24, "C": 24, "D": 18, "E": 18, "F": 18, "G": 24, "H": 18, "I": 24, "J": 18}
    for col, width in widths.items():
        ws_o.column_dimensions[col].width = width
    for row in ws_o.iter_rows(min_row=2, max_row=1000, min_col=1, max_col=10):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center")

    dv_grade = DataValidation(type="list", formula1='"A高意向,B意向,C潜在,D沉睡"', allow_blank=True)
    dv_model = DataValidation(type="list", formula1='"Deepal S07,Deepal S05,Deepal SL03,ROX ADAMAS"', allow_blank=True)
    ws_o.add_data_validation(dv_grade)
    ws_o.add_data_validation(dv_model)
    dv_grade.add("H2:H1000")
    dv_model.add("G2:G1000")

    ws_t = wb[SHEET_TEMPLATE]
    labels = ["客户公司：", "对接人：", "目标车型：", "国家：", "客户等级："]
    for idx, label in enumerate(labels, 1):
        ws_t[f"A{idx}"] = label
    record_headers = ["跟进日期", "跟进渠道", "跟进内容详情", "客户反馈", "下次跟进计划"]
    for idx, header in enumerate(record_headers, 1):
        ws_t.cell(7, idx, header)
    for row in ws_t.iter_rows(min_row=1, max_row=5, min_col=1, max_col=6):
        for cell in row:
            cell.font = Font(bold=True)
            cell.fill = pale
            cell.border = border
            cell.alignment = Alignment(vertical="center")
    for cell in ws_t[7][:5]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = gold
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws_t.iter_rows(min_row=8, max_row=1000, min_col=1, max_col=5):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws_t.row_dimensions[7].height = 26
    ws_t.column_dimensions["A"].width = 20
    ws_t.column_dimensions["B"].width = 20
    ws_t.column_dimensions["C"].width = 34
    ws_t.column_dimensions["D"].width = 34
    ws_t.column_dimensions["E"].width = 34
    dv_channel = DataValidation(type="list", formula1='"WhatsApp,邮件,电话,IG/Facebook社媒"', allow_blank=True)
    ws_t.add_data_validation(dv_channel)
    dv_channel.add("B8:B1000")


def main() -> None:
    wb = reset_workbook()
    style_workbook(wb)
    wb.save(OUT_XLSX)
    OUT_BAS.write_text(VBA_CODE, encoding="utf-8-sig")
    print(OUT_XLSX)
    print(OUT_BAS)


if __name__ == "__main__":
    main()
