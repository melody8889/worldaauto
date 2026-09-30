from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
OUTPUT = BASE / "Russia_Local_Leads_Rescreened.xlsx"

target_leads = [
    {
        "name": "Drom - Li/Lixiang 在售广告卖家池",
        "type": "俄罗斯本地车商挖掘入口",
        "city": "俄罗斯多城市",
        "why": "不是品牌官方站，而是车型广告池；可逐条找到正在卖 Li/Lixiang 的俄罗斯本地车商或中间商。",
        "not_china_official": "平台型入口，本身不是中国品牌海外公司；需要进入每条广告识别卖家。",
        "contact": "平台内电话/消息，逐条广告查看",
        "domestic_way": "国内可直接打开网页筛账号；Telegram 不作为必需渠道。",
        "source": "https://auto.drom.ru/li/ ; https://auto.drom.ru/li/l9/",
        "priority": "高",
        "opening": "Здравствуйте, увидел ваше объявление по Li L7/L9 на Drom. Есть свежие предложения из Китая, могу отправить прайс и срок поставки.",
        "next_step": "按城市筛莫斯科、叶卡捷琳堡、新西伯利亚、海参崴；优先同账号多台车源。",
    },
    {
        "name": "Avito - 中国车现车/代购广告卖家池",
        "type": "俄罗斯本地中小车商入口",
        "city": "俄罗斯多城市",
        "why": "Avito 上大量中小车商发布中国车广告，目标更接近真实买家/中间商。",
        "not_china_official": "平台入口；不要联系平台，只联系广告页里的本地卖家账号。",
        "contact": "平台内电话/消息，需登录或点开广告",
        "domestic_way": "国内可先筛选账号和车型；后续可让俄语兼职代发平台消息。",
        "source": "https://www.avito.ru/all/avtomobili",
        "priority": "高",
        "opening": "Здравствуйте, увидел ваше объявление по китайскому авто. Можем предложить поставки из Китая по Li/Zeekr/Jetour/Tank.",
        "next_step": "优先找多车源商家、店铺号、长期更新账号；跳过单台私人卖家。",
    },
    {
        "name": "Auto.ru - Авто на заказ / 中国品牌广告卖家池",
        "type": "俄罗斯经销商/车商入口",
        "city": "俄罗斯多城市",
        "why": "平台有“Авто на заказ”和经销商广告，适合按 Jetour、Tank、Changan、Geely、Li、Zeekr 找具体车商。",
        "not_china_official": "平台型入口；需要看广告卖家是否多品牌、本地地址、是否非官方单品牌。",
        "contact": "平台内电话/表单/消息",
        "domestic_way": "国内可先做筛选；联系可用平台表单、电话或俄语兼职。",
        "source": "https://auto.ru/",
        "priority": "高",
        "opening": "Здравствуйте, увидел у вас авто на заказ/китайские авто. Можем дать экспортные цены из Китая по конкретным моделям.",
        "next_step": "按品牌和城市找具体商家账号，不联系 Auto.ru 官方邮箱。",
    },
    {
        "name": "Auto.kolesa.ru - 经销商车源目录",
        "type": "俄罗斯本地经销商线索入口",
        "city": "俄罗斯多城市",
        "why": "聚合新车/二手车和经销商报价，可挖二线城市本地经销商。",
        "not_china_official": "目录入口；具体商家需二次判断，排除官方单品牌经销。",
        "contact": "广告页/经销商页展示",
        "domestic_way": "国内可先筛商家名和官网，再用电话/表单联系。",
        "source": "https://auto.kolesa.ru/all-auto",
        "priority": "中",
        "opening": "Здравствуйте, работаем с экспортом авто из Китая. Можем предложить поставки под ваш спрос по популярным китайским моделям.",
        "next_step": "重点找二线城市、多品牌、非官方单一品牌店。",
    },
    {
        "name": "Drom - 中国品牌非 Li 车型广告池",
        "type": "多品牌车商挖掘入口",
        "city": "俄罗斯多城市",
        "why": "可按 Jetour T2、Tank 300/500、Zeekr、Voyah、Changan、Geely 等车型挖本地卖家。",
        "not_china_official": "车型广告入口；不等于品牌官方渠道。",
        "contact": "平台内展示",
        "domestic_way": "国内先筛账号；找俄语兼职代发首轮消息更稳。",
        "source": "https://auto.drom.ru/",
        "priority": "高",
        "opening": "Здравствуйте, видел ваш автомобиль [车型] на Drom. Можем дать цену из Китая и срок поставки по похожим комплектациям.",
        "next_step": "每天按 3 个热门车型各挖 10 个账号，建立滚动开发池。",
    },
]

referral_channels = [
    {
        "name": "中俄汽车物流/清关代理",
        "role": "转介绍渠道",
        "why": "他们每天接触真实进口需求，且你在国内更容易通过中文端物流公司联系。",
        "how": "找中俄汽运、铁路、满洲里/绥芬河/霍尔果斯、海参崴清关代理。",
        "contact": "官网电话、微信、国内业务员",
        "opening": "你们有俄罗斯客户需要中国车源报价吗？成交可以给介绍费，报价和车源我这边负责。",
        "priority": "高",
    },
    {
        "name": "俄罗斯中国车维修/配件店",
        "role": "转介绍渠道",
        "why": "维修店知道哪些车商卖得多、谁缺车、谁缺配件和售后资源。",
        "how": "搜索 сервис китайских автомобилей / ремонт Li Auto / ремонт Zeekr / запчасти китайских автомобилей。",
        "contact": "官网电话、VK、地图电话",
        "opening": "你们有没有合作车商需要中国车源？我可以提供现车报价和配件支持信息。",
        "priority": "中高",
    },
    {
        "name": "俄语本地兼职/代开发员",
        "role": "执行渠道",
        "why": "解决你在国内打不开 Telegram、平台消息不方便、俄语沟通慢的问题。",
        "how": "找俄罗斯本地留学生、华人、俄语兼职，按有效线索付费。",
        "contact": "微信/国内招聘/熟人介绍",
        "opening": "帮我每天从 Drom/Avito/Auto.ru 找 30 个多车源中国车商，并完成第一轮俄语沟通。",
        "priority": "高",
    },
]

excluded = [
    {
        "name": "Li Auto Russia official distributor",
        "reason": "品牌官方/准官方俄语站，可能是中国品牌海外体系或授权体系，不适合作为主攻中小客户。",
        "use": "只作为市场观察和下游 dealer 入口。",
        "source": "https://liautoofficial.ru/",
    },
    {
        "name": "Li Motors / Lixiang Motors",
        "reason": "明显围绕单一中国品牌 Li/Lixiang 体系，有“заказать Lixiang из Китая / стать партнером”等官方或准官方特征。",
        "use": "可观察价格和渠道，不作为优先客户。",
        "source": "https://li-motors.ru/",
    },
    {
        "name": "Оками Li Auto / Okami Li Auto",
        "reason": "官方经销商属性强，可能已有固定供货/授权体系，不像开放型中小进口商。",
        "use": "降级为观察对象；除非你要找区域合作，不放主攻。",
        "source": "https://okami-liauto.ru/",
    },
    {
        "name": "T-Auto catalog",
        "reason": "金融/大平台销售入口，不是中小客户，直接开发概率低。",
        "use": "只看车型和价格趋势。",
        "source": "https://www.tbank.ru/auto/catalog/",
    },
]

rules = [
    ["排除", "官网写 официальный дилер / официальный импортер / distributor，且只围绕单一中国品牌", "大概率官方/准官方体系，不主攻"],
    ["排除", "公司名、页面、社媒明显是中国品牌海外站或中国团队运营", "可能是竞争渠道或封闭授权体系"],
    ["保留", "平台广告账号长期发布多台不同品牌中国车", "更像本地车商/中间商"],
    ["保留", "俄罗斯本地地址、多个品牌、二线城市、可电话联系", "更符合中小客户画像"],
    ["保留", "物流清关、维修配件、改装店", "作为转介绍渠道，不一定直接买车"],
    ["谨慎", "只有 Telegram，无官网电话/邮箱/平台消息", "你国内操作困难，需俄语兼职代联"],
]


def add_table(sheet, display_name, style_name):
    end_col = get_column_letter(sheet.max_column)
    table = Table(displayName=display_name, ref=f"A1:{end_col}{sheet.max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name=style_name,
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def style_sheet(sheet, widths, fill_color):
    header_fill = PatternFill("solid", fgColor=fill_color)
    white = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
    body = Font(name="Microsoft YaHei", size=10)
    side = Side(style="thin", color="D9E2F3")
    sheet.freeze_panes = "A2"
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(top=side, bottom=side, left=side, right=side)
            if cell.row == 1:
                cell.fill = header_fill
                cell.font = white
            else:
                cell.font = body
    for idx, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(idx)].width = width
    for row_num in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row_num].height = 78


wb = Workbook()

ws = wb.active
ws.title = "主攻名单"
headers = [
    "名称",
    "类型",
    "城市/区域",
    "为什么保留",
    "非中国官方判断",
    "公开联系方式",
    "国内可操作方式",
    "来源链接",
    "优先级",
    "俄语开场白",
    "下一步",
]
ws.append(headers)
for row in target_leads:
    ws.append([
        row["name"],
        row["type"],
        row["city"],
        row["why"],
        row["not_china_official"],
        row["contact"],
        row["domestic_way"],
        row["source"],
        row["priority"],
        row["opening"],
        row["next_step"],
    ])
style_sheet(ws, [30, 24, 18, 44, 42, 30, 36, 48, 10, 58, 42], "1F4E79")
add_table(ws, "RescreenedTargetLeads", "TableStyleMedium2")

ref = wb.create_sheet("转介绍渠道")
ref.append(["名称", "角色", "为什么有用", "怎么找", "联系方式", "开场方向", "优先级"])
for row in referral_channels:
    ref.append([row["name"], row["role"], row["why"], row["how"], row["contact"], row["opening"], row["priority"]])
style_sheet(ref, [30, 18, 44, 48, 28, 50, 10], "6B7C3A")
add_table(ref, "ReferralChannels", "TableStyleMedium4")

ex = wb.create_sheet("排除观察")
ex.append(["名称", "排除原因", "还可以怎么用", "来源链接"])
for row in excluded:
    ex.append([row["name"], row["reason"], row["use"], row["source"]])
style_sheet(ex, [34, 58, 42, 50], "7F1D1D")
add_table(ex, "ExcludedWatchlist", "TableStyleMedium9")

rule = wb.create_sheet("筛选规则")
rule.append(["分类", "判断标准", "处理方式"])
for row in rules:
    rule.append(row)
style_sheet(rule, [14, 62, 44], "4B5563")
add_table(rule, "ScreeningRules", "TableStyleMedium1")

note = wb.create_sheet("说明")
note.append(["项目", "内容"])
note_rows = [
    ["生成日期", "2026-07-14"],
    ["本版变化", "已将 Li Auto 官方/准官方体系、疑似中国品牌海外渠道、大平台销售入口从主攻名单移到排除观察。"],
    ["核心思路", "主攻正在平台卖中国车的俄罗斯本地车商账号，以及物流清关、维修配件等转介绍渠道。"],
    ["Telegram 限制", "因国内无法稳定使用 Telegram，本版优先标注官网、平台内消息、电话、表单、俄语兼职代开发等可操作路径。"],
    ["下一步建议", "不要再只做公司名单。每天从 Drom/Avito/Auto.ru 按车型挖具体广告账号，记录账号名、城市、车型、电话/表单、是否多车源。"],
    ["合规提醒", "俄罗斯方向需要筛查客户主体、最终用途、付款路径和出口限制，避免不透明交易。"],
]
for row in note_rows:
    note.append(row)
style_sheet(note, [20, 100], "0F766E")

wb.save(OUTPUT)

check = load_workbook(OUTPUT)
main = check["主攻名单"]
assert main.max_row == len(target_leads) + 1
assert main["A1"].value == "名称"
assert main["A2"].value.startswith("Drom")
print(ascii({
    "file": str(OUTPUT),
    "sheets": check.sheetnames,
    "main_rows": main.max_row,
    "first_row": [main.cell(2, i).value for i in range(1, 5)],
}))
