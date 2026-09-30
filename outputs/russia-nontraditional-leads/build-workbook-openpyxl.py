from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

output = r"C:\Users\Melody\Documents\汽车出口独立站\outputs\russia-nontraditional-leads\俄罗斯非常规客户开发清单.xlsx"

direct_leads = [
    {
        "name": "Оками Li Auto / Okami Li Auto",
        "type": "区域经销商/中国新能源车",
        "city": "Екатеринбург",
        "channel": "车商/代购",
        "priority": "高",
        "fit": "官网明确为 Li Auto/Lixiang 叶卡捷琳堡经销商，有现车、试驾、企业客户入口，适合用中国现车/补充车源切入。",
        "public_contact": "+7 (343) 379-62-26",
        "social": "未公开",
        "source": "https://okami-liauto.ru/",
        "suggested_open": "Здравствуйте, увидел у вас Li Auto/Lixiang в Екатеринбурге. Можем быстро дать экспортную цену по L6/L7/L9 из Китая и расчёт поставки до РФ.",
        "note": "直接联系销售/采购/企业客户部门；先不要讲公司介绍，先给车型+价格表。",
    },
    {
        "name": "Li Motors / Lixiang Motors",
        "type": "Li/Lixiang 车源与经销网络",
        "city": "俄罗斯多城市/官网未集中列明",
        "channel": "车商/代购",
        "priority": "高",
        "fit": "页面出现“Заказать Lixiang из Китая”“Авто в наличии”“Стать партнером”等入口，适合找合作/供货/区域代理负责人。",
        "public_contact": "+7 (343) 302-20-67（叶卡捷琳堡经销页）",
        "social": "Telegram: https://t.me/lixiangmotors",
        "source": "https://li-motors.ru/ ; https://li-motors.ru/dealers/ekaterinburg",
        "suggested_open": "Здравствуйте. Видим, что вы работаете с Lixiang и заказом из Китая. Можем предложить стабильные экспортные предложения по L6/L7/L9 с фото, VIN и сроком отгрузки.",
        "note": "优先问采购/партнерство，不要只找零售销售。",
    },
    {
        "name": "Li Auto Russia official distributor",
        "type": "品牌俄语站/经销入口",
        "city": "俄罗斯",
        "channel": "车商/代购",
        "priority": "中",
        "fit": "官网有“Стать дилером”“Корпоративные продажи”“Найти дилера”，适合继续挖其下游经销商和企业销售入口。",
        "public_contact": "未公开",
        "social": "Telegram: https://t.me/liautoofficial ; VK: https://vk.com/liautoofficial",
        "source": "https://liautoofficial.ru/",
        "suggested_open": "Здравствуйте. Мы поставляем автомобили из Китая и хотим обсудить B2B-поставки/партнёрство по Li Auto для дилеров в РФ.",
        "note": "更像入口，不一定是中小客户；用来找下游 dealer 列表。",
    },
    {
        "name": "Drom Li / Li L9 listings",
        "type": "车型广告平台/可挖具体卖家",
        "city": "俄罗斯多城市",
        "channel": "车商/代购",
        "priority": "高",
        "fit": "Li/Lixiang 有大量在售广告；每条广告背后是正在卖中国新能源车的车商或个人中间商，比公开进口商名单更活跃。",
        "public_contact": "平台内展示，需逐条进入广告查看",
        "social": "Telegram: https://t.me/drom ; VK: https://vk.com/drom/",
        "source": "https://auto.drom.ru/li/ ; https://auto.drom.ru/li/l9/",
        "suggested_open": "Здравствуйте, увидел ваше объявление по Li L7/L9 на Drom. Есть регулярные предложения из Китая, могу отправить свежий прайс и сроки поставки.",
        "note": "按城市筛：莫斯科、叶卡捷琳堡、新西伯利亚、海参崴；优先联系同一账号多台车源者。",
    },
    {
        "name": "Auto.ru - Авто на заказ / Chinese brands",
        "type": "车源平台/经销商入口",
        "city": "俄罗斯多城市",
        "channel": "车商/代购",
        "priority": "中",
        "fit": "平台有“Авто на заказ”“Для бизнеса”“Дилеры”，可按 Jetour/Changan/Geely/Tank/Li/Zeekr 找具体商家。",
        "public_contact": "平台内展示，需逐条进入广告查看",
        "social": "未公开",
        "source": "https://auto.ru/",
        "suggested_open": "Здравствуйте, увидел у вас китайские авто/авто на заказ. Можем дать экспортные цены из Китая по конкретным моделям и расчёт доставки.",
        "note": "不要联系平台邮箱；要联系广告页内车商账号。",
    },
    {
        "name": "Avito - автомобили / авто из Китая sellers",
        "type": "分类广告/中小车商入口",
        "city": "俄罗斯多城市",
        "channel": "车商/代购",
        "priority": "高",
        "fit": "大量小车商在 Avito 发中国车现车/代购广告，适合按车型直接打招呼。",
        "public_contact": "平台内展示，需登录/点开广告查看",
        "social": "未公开",
        "source": "https://www.avito.ru/all/avtomobili",
        "suggested_open": "Здравствуйте, увидел ваше объявление по китайскому авто. Есть поставки из Китая по Li/Zeekr/Jetour/Tank, могу отправить цены по 3-5 ходовым моделям.",
        "note": "优先找多车源商家，不要找单台私人卖家。",
    },
    {
        "name": "Auto.kolesa.ru dealer catalog",
        "type": "经销商车源平台",
        "city": "俄罗斯多城市",
        "channel": "车商/代购",
        "priority": "中",
        "fit": "搜索结果显示平台聚合官方经销商和新车/二手车，适合找二线城市经销商。",
        "public_contact": "平台内展示，需按车型/城市进入",
        "social": "未公开",
        "source": "https://auto.kolesa.ru/all-auto",
        "suggested_open": "Здравствуйте, работаем с экспортом авто из Китая. Можем предложить поставки под ваш спрос по популярным китайским моделям.",
        "note": "作为补充入口，不作为单独客户。",
    },
    {
        "name": "T-Auto catalog",
        "type": "金融/平台销售入口",
        "city": "俄罗斯",
        "channel": "车商/代购",
        "priority": "低",
        "fit": "平台有新车购买、申请和配送流程；可作为观察俄市场车型和价格的入口，直接成交可能性低。",
        "public_contact": "平台表单",
        "social": "未公开",
        "source": "https://www.tbank.ru/auto/catalog/",
        "suggested_open": "不建议作为第一批开发对象；用于看价格和车型趋势。",
        "note": "大平台，回复概率低。",
    },
]

playbook = [
    ["Drom/Avito/Auto.ru 车型广告", "按 Li L7/L9、Zeekr 001/009、Jetour T2、Tank 300/500、Changan、Geely 筛选", "优先联系同账号多台车的卖家", "开场只说看到某车型广告+能给中国出口价"],
    ["Li/Lixiang 经销体系", "从 liautoofficial、li-motors、okami 页面进入 dealer/partner/contact", "找采购、企业客户、合作部门", "不要问是否买车，问是否需要稳定中国车源/补充配置"],
    ["物流清关转介绍", "搜索俄语：доставка авто из Китая / растаможка авто из Китая / таможенный брокер авто", "问他们是否有客户需要中国现车报价", "给介绍佣金或一起赚服务费"],
    ["售后配件/维修店", "搜索俄语：сервис китайских автомобилей / ремонт Li Auto / ремонт Zeekr", "维修店知道谁卖得多、谁缺货", "不要推销车，先问是否认识车商需要稳定车源"],
    ["区域打法", "叶卡捷琳堡、新西伯利亚、喀山、乌法、克拉斯诺达尔、海参崴", "二线城市比莫斯科更缺中国供应链", "用具体到城的到货时间和成本表切入"],
]

headers = ["名称", "类型", "城市/区域", "渠道", "优先级", "为什么适合开发", "公开联系方式", "Telegram/VK/社媒", "来源链接", "建议开场白", "备注"]
keys = ["name", "type", "city", "channel", "priority", "fit", "public_contact", "social", "source", "suggested_open", "note"]

wb = Workbook()
ws = wb.active
ws.title = "客户与渠道清单"
ws.append(headers)
for lead in direct_leads:
    ws.append([lead[k] for k in keys])

play = wb.create_sheet("开发打法")
play.append(["路径", "怎么找", "筛选标准", "第一句话方向"])
for row in playbook:
    play.append(row)

note = wb.create_sheet("说明")
for row in [
    ["生成日期", "2026-07-14"],
    ["原则", "只记录网页公开可验证的联系方式；未能稳定提取的联系方式标为未公开或平台内展示。"],
    ["重要提醒", "俄罗斯方向开发需做合规审查，避免受限制主体、受限制车型/用途、异常付款和不透明最终买家。"],
    ["使用方法", "第一批优先联系高优先级直接对象；同时每天从 Drom/Avito/Auto.ru 按热门车型挖 20 个正在卖车的具体账号。"],
    ["推荐车型切入", "Li L6/L7/L9、Zeekr 001/009、Jetour T2、Tank 300/500、Changan、Geely、Voyah。"],
    ["不要这样做", "不要群发公司介绍；不要只发“we export cars”。每条消息必须带对方正在卖的车型。"],
    ["后续补充", "若需要更大名单，建议按城市+车型逐条进入平台广告人工提取车商账号。"],
    ["来源", "见客户与渠道清单的来源链接列。"],
]:
    note.append(row)

header_fill = PatternFill("solid", fgColor="1F4E79")
play_fill = PatternFill("solid", fgColor="6B7C3A")
note_fill = PatternFill("solid", fgColor="F2F2F2")
white_font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
body_font = Font(name="Microsoft YaHei", size=10)
thin_blue = Side(style="thin", color="D9E2F3")
thin_green = Side(style="thin", color="E2E8D5")

def format_sheet(sheet, widths, header_color):
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        cell.fill = header_color
        cell.font = white_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in sheet.iter_rows():
        for cell in row:
            if cell.row != 1:
                cell.font = body_font
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(top=thin_blue, bottom=thin_blue, left=thin_blue, right=thin_blue)
    for i, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    for row in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row].height = 72

format_sheet(ws, [28, 22, 20, 12, 10, 42, 34, 38, 48, 58, 36], header_fill)
tab = Table(displayName="RussiaNontraditionalLeads", ref=f"A1:K{ws.max_row}")
tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False, showLastColumn=False)
ws.add_table(tab)

format_sheet(play, [28, 50, 40, 50], play_fill)
tab2 = Table(displayName="RussiaLeadPlaybook", ref=f"A1:D{play.max_row}")
tab2.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True, showFirstColumn=False, showLastColumn=False)
play.add_table(tab2)

for row in note.iter_rows():
    for cell in row:
        cell.font = body_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")
for cell in note["A"]:
    cell.fill = note_fill
    cell.font = Font(name="Microsoft YaHei", size=10, bold=True)
note.column_dimensions["A"].width = 18
note.column_dimensions["B"].width = 95

wb.save(output)

check = load_workbook(output)
ws_check = check["客户与渠道清单"]
print(check.sheetnames)
print([ws_check.cell(1, i).value for i in range(1, 6)])
print([ws_check.cell(2, i).value for i in range(1, 6)])
