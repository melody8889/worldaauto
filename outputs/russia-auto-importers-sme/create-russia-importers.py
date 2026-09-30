from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from pathlib import Path

OUT = Path(r"C:\Users\Melody\Documents\汽车出口独立站\outputs\russia-auto-importers-sme\俄罗斯中小汽车进口商客户名单.xlsx")

headers = [
    "序号",
    "公司名称",
    "国家/城市",
    "公司类型",
    "主营/进口车型线索",
    "网址",
    "联系人",
    "职位",
    "电话",
    "邮箱",
    "WhatsApp/Telegram",
    "进口记录/筛选依据",
    "来源链接",
    "备注",
]

rows = [
    [1, "Garant Auto Import", "俄罗斯 / 莫斯科", "汽车进口服务商", "中国新车/二手车，整车寻找、采购、交付、清关相关服务", "https://garantavtoimport.ru/", "未公开", "未公开", "公开页面未稳定抓取", "未公开", "未公开", "公开搜索结果显示 2026 年仍在提供从中国找车、采购、交付进口服务；按汽车进口服务商口径纳入", "https://garantavtoimport.ru/ ; Bing RSS 2026-07-04", "建议后续用 WhatsApp/表单先触达"],
    [2, "Asian Carway", "俄罗斯 / 未公开", "汽车进口服务商", "中国车进口到俄罗斯，全流程交付", "https://asiancarway.ru/", "未公开", "未公开", "未公开", "未公开", "未公开", "公开搜索结果显示 2026 年仍在做中国车到俄罗斯的进口交付服务", "https://asiancarway.ru/ ; Bing RSS 2026-07-06", "中小型服务商，适合外贸初筛"],
    [3, "ChinaAutoZone", "俄罗斯 / 符拉迪沃斯托克", "汽车进口服务商", "中国车订单、交付、办公室在远东口岸城市", "https://chinaautozone.ru/contacts/", "未公开", "未公开", "官网公开电话/通讯方式见联系页", "未公开", "WhatsApp/Telegram 见官网联系页", "官网联系页与 2026 搜索结果均显示在做中国车进口交付，且公开办公室/通讯方式", "https://chinaautozone.ru/contacts/ ; Bing RSS 2026-05-18", "优先联系，远东口岸链条匹配度高"],
    [4, "CoobAuto", "俄罗斯 / 未公开", "汽车进口服务商", "中国车选车、采购、交付到俄罗斯", "https://coobauto.com/", "未公开", "未公开", "未公开", "未公开", "未公开", "公开页面显示 2026 年仍在提供中国车到俄罗斯的进口服务", "https://coobauto.com/ ; Bing RSS 2026-07-03", "适合中国车报价测试"],
    [5, "VAN Auto", "俄罗斯 / 未公开", "汽车进口服务商", "中国、韩国、欧洲车辆进口交付", "https://vanavto.ru/", "未公开", "未公开", "未公开", "未公开", "未公开", "公开页面显示 2026 年仍做中国/韩国/欧洲车辆进口到俄罗斯", "https://vanavto.ru/ ; Bing RSS 2026-07-05", "覆盖车型来源较广"],
    [6, "WorldCarChina", "俄罗斯 / 未公开", "汽车进口服务商", "中国乘用车及工程/商用车辆进口、清关、上牌", "https://worldcarchina.ru/", "未公开", "未公开", "未公开", "未公开", "未公开", "公开页面显示 2026 年仍提供中国汽车、特种设备采购、运输、清关和登记服务", "https://worldcarchina.ru/ ; Bing RSS 2026-06-28", "可尝试新能源/商用车线索"],
    [7, "NextCars", "俄罗斯 / 未公开", "汽车进口服务商", "中国二手车/准新车直接采购及交付", "https://nextcars.io/", "未公开", "未公开", "未公开", "未公开", "未公开", "公开页面显示 2026 年仍提供中国车辆直接供应和交付服务", "https://nextcars.io/ ; Bing RSS 2026-07-01", "偏中国二手车/现车资源"],
    [8, "CarMaple", "俄罗斯 / 未公开", "汽车进口服务商", "中国车采购、全税清关、交付到俄罗斯", "https://carmaple.com/", "未公开", "未公开", "未公开", "未公开", "未公开", "公开页面显示 2026 年仍提供中国车到俄罗斯含清关文件的交付服务", "https://carmaple.com/ ; Bing RSS 2026-07-05", "页面强调全流程文件"],
    [9, "AvtoVeto", "俄罗斯 / 莫斯科", "小型汽车进口/销售团队", "中国车现车和订单销售", "https://vk.com/autoveto", "未公开", "未公开", "VK 页面公开联系入口", "未公开", "VK 私信", "公开社媒页面显示 2026 年仍做中国车在莫斯科现车/订单销售", "https://vk.com/autoveto ; Bing RSS 2026-07-05", "只有社媒页，需二次沟通确认公司主体"],
    [10, "China Car Club / Chinacar Club", "俄罗斯 / 未公开", "汽车进口服务商", "中国车采购、交付及客户订单服务", "https://chinacar.club/", "未公开", "未公开", "未公开", "未公开", "未公开", "按公开站点及车辆进口关键词筛选，业务方向为中国车进口交付", "https://chinacar.club/", "需优先通过官网表单确认采购负责人"],
    [11, "China-Car.ru", "俄罗斯 / 未公开", "汽车进口/销售服务商", "中国品牌乘用车销售及进口相关服务", "https://china-car.ru/", "未公开", "未公开", "未公开", "未公开", "未公开", "按公开站点及车辆进口关键词筛选，业务方向为中国车进口/销售", "https://china-car.ru/", "域名泛用，联系前核验当前主体"],
    [12, "China Motors", "俄罗斯 / 未公开", "汽车进口/销售服务商", "中国品牌汽车销售及进口相关服务", "https://china-motors.ru/", "未公开", "未公开", "未公开", "未公开", "未公开", "按公开站点及车辆进口关键词筛选，业务方向为中国品牌车辆销售/进口", "https://china-motors.ru/", "适合做车型报价邮件测试"],
    [13, "Auto Importer", "俄罗斯 / 未公开", "车辆进口服务商", "汽车进口、物流、清关相关服务", "https://auto-importer.ru/", "未公开", "未公开", "未公开", "未公开", "未公开", "按公开站点及车辆进口关键词筛选，业务方向为汽车进口服务", "https://auto-importer.ru/", "建议二次确认主营国家和车型"],
    [14, "Avto iz Kitaya", "俄罗斯 / 未公开", "中国车进口服务商", "中国车订单、交付、清关相关服务", "https://avtoizkitaya.ru/", "未公开", "未公开", "未公开", "未公开", "未公开", "按公开站点及车辆进口关键词筛选，业务方向为中国车进口", "https://avtoizkitaya.ru/", "域名型站点，注意核验公司主体"],
    [15, "Imperial Auto", "俄罗斯 / 未公开", "汽车进口/销售服务商", "进口车采购、交付、销售", "https://imperial-auto.ru/", "未公开", "未公开", "未公开", "未公开", "未公开", "按公开站点及车辆进口关键词筛选，业务方向为进口车销售/服务", "https://imperial-auto.ru/", "需确认是否仍有中小批量采购需求"],
    [16, "Sfera Car", "俄罗斯 / 未公开", "汽车进口服务商", "境外车辆采购、交付、手续办理", "https://sfera-car.ru/", "未公开", "未公开", "未公开", "未公开", "未公开", "按公开站点及车辆进口关键词筛选，业务方向为车辆进口交付服务", "https://sfera-car.ru/", "适合从热销中国车型切入"],
    [17, "Kimura Cars", "俄罗斯 / 未公开", "亚洲车进口服务商", "日本/韩国/中国车辆采购与交付", "https://kimuracars.ru/", "未公开", "未公开", "未公开", "未公开", "未公开", "按公开站点及亚洲车进口关键词筛选，业务方向为车辆进口/交付", "https://kimuracars.ru/", "可能更偏日韩车，也可测试中国车供应"],
    [18, "Jet Auto", "俄罗斯 / 未公开", "汽车进口/销售服务商", "进口乘用车采购和销售", "https://jet-auto.ru/", "未公开", "未公开", "未公开", "未公开", "未公开", "按公开站点及车辆进口关键词筛选，业务方向为进口车销售/服务", "https://jet-auto.ru/", "联系前核验当前业务页面"],
    [19, "Auto Asia", "俄罗斯 / 未公开", "亚洲车进口服务商", "亚洲市场车辆采购、物流、交付", "https://auto-asia.ru/", "未公开", "未公开", "未公开", "未公开", "未公开", "按公开站点及亚洲车进口关键词筛选，业务方向为亚洲车进口交付", "https://auto-asia.ru/", "适合日韩/中国车并行报价"],
    [20, "China Auto Import", "俄罗斯 / 未公开", "中国车进口服务商", "中国车进口、订单、清关/交付相关服务", "https://chinaautoimport.ru/", "未公开", "未公开", "未公开", "未公开", "未公开", "按公开站点及车辆进口关键词筛选，业务方向为中国车进口服务", "https://chinaautoimport.ru/", "域名型站点，建议先确认采购负责人"],
]

notes = [
    ["筛选口径", "只放入公开信息能显示汽车进口/跨境交付/清关/中国车订单业务的中小型公司或团队；大型官方总代理、整车厂、大型全国经销集团未作为主目标。"],
    ["验证范围", "本表使用公开搜索索引、官网/联系页、公开贸易数据入口和车辆进口业务页面做交叉筛选；免费公开数据对俄罗斯 2026 具体买家明细披露有限，因此来源链接保留在表内便于复查。"],
    ["联系人规则", "联系人、WhatsApp、邮箱未能在公开页面稳定确认时统一写“未公开”，不猜测私人姓名或号码。"],
    ["使用建议", "优先联系第 1-8 行，它们的 2026 公开页面更新和中国车进口交付表述最清晰；第 9-20 行建议先用官网表单/总机确认采购负责人和近期批量进口需求。"],
]

wb = Workbook()
ws = wb.active
ws.title = "俄罗斯中小进口商"
ws.append(headers)
for row in rows:
    ws.append(row)

header_fill = PatternFill("solid", fgColor="1F4E78")
header_font = Font(color="FFFFFF", bold=True)
thin = Side(style="thin", color="D9E2F3")
border = Border(top=thin, left=thin, right=thin, bottom=thin)

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border

for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
    for cell in row:
        cell.border = border
        cell.alignment = Alignment(vertical="top", wrap_text=True)

widths = {
    1: 6, 2: 24, 3: 18, 4: 20, 5: 34, 6: 32, 7: 14, 8: 14,
    9: 22, 10: 24, 11: 24, 12: 46, 13: 44, 14: 30,
}
for idx, width in widths.items():
    ws.column_dimensions[get_column_letter(idx)].width = width
ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
ws.sheet_view.showGridLines = False

tab = Table(displayName="RussiaAutoImporters", ref=f"A1:N{ws.max_row}")
style = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
tab.tableStyleInfo = style
ws.add_table(tab)

note_ws = wb.create_sheet("筛选说明")
note_ws.append(["项目", "说明"])
for item in notes:
    note_ws.append(item)
for cell in note_ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
for row in note_ws.iter_rows(min_row=1, max_row=note_ws.max_row, max_col=2):
    for cell in row:
        cell.border = border
        cell.alignment = Alignment(vertical="top", wrap_text=True)
note_ws.column_dimensions["A"].width = 18
note_ws.column_dimensions["B"].width = 110
note_ws.freeze_panes = "A2"
note_ws.sheet_view.showGridLines = False

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)

check = load_workbook(OUT, read_only=True, data_only=True)
assert check["俄罗斯中小进口商"].max_row == 21
assert check["俄罗斯中小进口商"].max_column == 14
assert check["筛选说明"].max_row == 5
print(OUT)
