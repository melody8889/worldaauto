from pathlib import Path

from PIL import Image
import openpyxl


base = Path(r"D:\Worlda产品资料\待上传\2026-06-第五批产品上传")
xlsx = next(base.glob("*.xlsx"))

wb = openpyxl.load_workbook(xlsx, data_only=True)
print("WORKBOOK", xlsx.name)
for ws in wb.worksheets:
    print("SHEET", ws.title, ws.max_row, ws.max_column)
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
        if any(v is not None for v in row):
            print("\t".join("" if v is None else str(v) for v in row))

for folder in ["rox-adamas", "deepal-s05", "deepal-s07", "deepal-sl03"]:
    print("IMAGES", folder)
    for f in sorted((base / folder).iterdir()):
        if f.is_file():
            im = Image.open(f)
            print(f.name, im.size, im.mode)
